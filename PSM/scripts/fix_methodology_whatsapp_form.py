"""Fix WhatsApp / Google Form methodology wording in report fyp 1.docx."""
from docx import Document
from pathlib import Path

SOURCE = Path(r"c:\Users\Admin\Downloads\report fyp 1.docx")
OUTPUT = Path(r"c:\Users\Admin\Downloads\report fyp 1 - methodology fixed.docx")

REPLACEMENTS = [
    (
        "Primary data was collected through a structured Google Form questionnaire completed by the manager on 11 June 2026, while supplementary clarification was obtained through WhatsApp communication. The Google Form response and selected WhatsApp evidence are provided in Appendix B and Appendix C.",
        "Requirement gathering began with WhatsApp communication on 9 June 2026 to obtain consent and initial business context (Appendix C). A structured Google Form questionnaire was then completed by the manager on 11 June 2026 as the main structured primary data source (Appendix B). Together, these sources supported the requirement analysis in Chapters 2 and 3.",
    ),
    (
        "based on primary data collected through the Google Form response in Appendix B and supplementary WhatsApp communication in Appendix C.",
        "based on primary data collected through WhatsApp communication in Appendix C and the structured Google Form response in Appendix B.",
    ),
    (
        "Structured Google Form questionnaire completed by the Kak Norie manager on 11 June 2026 (Appendix B).",
        "WhatsApp communication with the Kak Norie manager (Nazhif) from 9 June 2026 onwards for consent, initial business context, and exploratory clarification (Appendix C).",
    ),
    (
        "Supplementary WhatsApp communication used for clarification and confirmation of business processes and requirements (Appendix C).",
        "Structured Google Form questionnaire completed on 11 June 2026 after the initial WhatsApp discussions (Appendix B).",
    ),
    (
        "The Google Form served as the main source of information regarding current recording practices, defect types, user roles, batch management, expiry control, corrective actions, and reporting requirements. The supplementary communication was used to clarify and confirm information obtained during requirement gathering.",
        "Requirement gathering followed the sequence WhatsApp first, then Google Form. WhatsApp established contact and informed the questionnaire scope. The Google Form served as the main structured primary data source for current recording practices, defect types, user roles, batch management, expiry control, corrective actions, and reporting requirements.",
    ),
    (
        "Requirement gathering began with the collection of primary data from Kak Norie. Literature review was then conducted to support requirement validation and evaluate alternative techniques.",
        "Requirement gathering began with WhatsApp communication with the Kak Norie manager, followed by the structured Google Form questionnaire. Literature review was then conducted to support requirement validation and evaluate alternative techniques.",
    ),
    (
        "This phase focused on understanding the current quality management practices at Kak Norie. Information was collected through a structured questionnaire and supplemented by follow-up communication. Topics included defect recording, batch management, expiry control, corrective actions, user roles, and reporting requirements.",
        "This phase focused on understanding the current quality management practices at Kak Norie. Information was collected first through WhatsApp communication for consent and initial clarification (9 June 2026), then through a structured Google Form questionnaire completed on 11 June 2026. Topics included defect recording, batch management, expiry control, corrective actions, user roles, and reporting requirements.",
    ),
    (
        "primary data collection through a structured questionnaire with supplementary WhatsApp communication.",
        "primary data collection beginning with WhatsApp communication (Appendix C) and continuing with a structured Google Form questionnaire (Appendix B).",
    ),
    (
        "describe the manual quality recording practices confirmed through the Google Form response in Appendix B.",
        "describe the manual quality recording practices confirmed through WhatsApp communication in Appendix C and the Google Form response in Appendix B.",
    ),
    (
        "Based on the interview conducted with the business owner, one of the most common defects is incorrect expiry date printing.",
        "Based on WhatsApp clarification and the Google Form response from the manager (Nazhif), one of the most common defects is incorrect expiry date printing.",
    ),
    (
        "The interview mentioned a possible batch format such as (product)-B-001.",
        "The Google Form response mentioned a possible batch format such as (product)-B-001.",
    ),
    (
        "The Google Form questionnaire serves as the main primary data source for the requirement analysis and current system investigation presented in Chapters 2 and 3 of this project. Additional clarification obtained through other communication channels is documented separately in Appendix C and should be considered supplementary to the primary questionnaire findings.",
        "The Google Form questionnaire serves as the main structured primary data source for the requirement analysis and current system investigation presented in Chapters 2 and 3 of this project. Initial requirement gathering contact and exploratory clarification through WhatsApp are documented in Appendix C (9 June 2026). The Google Form was administered after those discussions to collect comprehensive structured responses on 11 June 2026.",
    ),
    (
        "A structured Google Form questionnaire was developed to collect primary data from the management of Kak Norie / Retort Niaga regarding the current quality defect management process.",
        "Following initial WhatsApp discussions with the manager (Appendix C), a structured Google Form questionnaire was developed to collect primary data from the management of Kak Norie / Retort Niaga regarding the current quality defect management process.",
    ),
    (
        "Kak Norie Google Form interview response",
        "Kak Norie Google Form questionnaire response",
    ),
    (
        "PSM appendices such as Gantt chart, interview evidence, and system screenshots",
        "PSM appendices such as Gantt chart, requirement gathering evidence, and system screenshots",
    ),
    (
        "Requirement gathering through Google Form and WhatsApp communication",
        "Requirement gathering (WhatsApp + Google Form)",
    ),
    (
        "Google Form questionnaire, supplementary communication, review of business practices",
        "WhatsApp communication (initial contact, 9 June 2026); structured Google Form questionnaire (11 June 2026); review of business practices",
    ),
    (
        "Requirement list, primary data, business information",
        "WhatsApp evidence (Appendix C); Google Form CSV (Appendix B); initial requirement list",
    ),
    (
        "Consolidate findings from primary data and validate expected system functions",
        "Consolidate WhatsApp and Google Form findings and validate expected system functions",
    ),
    (
        "Quality defect tracking interview response (Google Form)",
        "Quality defect tracking questionnaire response (Google Form)",
    ),
]


def replace_in_paragraphs(doc: Document) -> int:
    count = 0
    for para in doc.paragraphs:
        original = para.text
        updated = original
        for old, new in REPLACEMENTS:
            if old in updated:
                updated = updated.replace(old, new)
        if updated != original:
            para.text = updated
            count += 1
    return count


def replace_in_tables(doc: Document) -> int:
    count = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                original = cell.text
                updated = original
                for old, new in REPLACEMENTS:
                    if old in updated:
                        updated = updated.replace(old, new)
                if updated != original:
                    cell.text = updated
                    count += 1
    return count


def main() -> None:
    doc = Document(SOURCE)
    para_count = replace_in_paragraphs(doc)
    table_count = replace_in_tables(doc)
    doc.save(OUTPUT)
    print(f"Updated {para_count} paragraph(s) and {table_count} table cell(s)")
    print(f"Saved: {OUTPUT}")
    print("Close report fyp 1.docx in Word, then replace it with this file if desired.")


if __name__ == "__main__":
    main()
