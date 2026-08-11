"""
Generate QDTS supervisor presentation (PPTX).
Run: python PSM/scripts/generate_presentation.py
Output: PSM/QDTS_Presentation_Slides.pptx
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "QDTS_Presentation_Slides.pptx"
ERD_SVG = ROOT / "diagrams" / "erd-qdts-core-report.svg"

# Brand colors
TEAL = RGBColor(0x17, 0x85, 0x6A)
TEAL_DARK = RGBColor(0x14, 0x63, 0x56)
INK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF4, 0xFA, 0xF7)


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title_text):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(9), Inches(0.7))
    tf = title.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = INK
    p.font.name = "Segoe UI"


def add_bullets(slide, items, top=1.2, size=18):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(8.8), Inches(5.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.name = "Segoe UI"
        p.font.color.rgb = INK
        p.space_after = Pt(10)


def add_footer(slide, text="QDTS · Kak Norie / Retort Niaga"):
    foot = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(9), Inches(0.35))
    p = foot.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.color.rgb = MUTED
    p.font.name = "Segoe UI"


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_BG)

    accent = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.18))
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()

    t1 = slide.shapes.add_textbox(Inches(0.6), Inches(2.0), Inches(8.8), Inches(1.2))
    p = t1.text_frame.paragraphs[0]
    p.text = "Quality Defect Tracking System (QDTS)"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = INK
    p.font.name = "Segoe UI"

    t2 = slide.shapes.add_textbox(Inches(0.6), Inches(3.2), Inches(8.8), Inches(0.6))
    p = t2.text_frame.paragraphs[0]
    p.text = "for Kak Norie / Retort Niaga"
    p.font.size = Pt(22)
    p.font.color.rgb = TEAL_DARK
    p.font.name = "Segoe UI"

    t3 = slide.shapes.add_textbox(Inches(0.6), Inches(4.5), Inches(8.8), Inches(1.5))
    tf = t3.text_frame
    for i, line in enumerate([
        "[Your Name] · [Matric Number]",
        "[Programme / Faculty]",
        "PSM1 Presentation · Supervisor Review",
    ]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(16)
        p.font.color.rgb = MUTED
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)


def content_slide(prs, title, bullets, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, title)
    add_bullets(slide, bullets)
    add_footer(slide, footer or "QDTS · Kak Norie / Retort Niaga · ~5 min slides + ~15 min demo")


def two_column_slide(prs, title, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, title)

    for col, (ctitle, citems, x) in enumerate([
        (left_title, left_items, 0.5),
        (right_title, right_items, 5.1),
    ]):
        ht = slide.shapes.add_textbox(Inches(x), Inches(1.15), Inches(4.3), Inches(0.4))
        p = ht.text_frame.paragraphs[0]
        p.text = ctitle
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEAL_DARK
        p.font.name = "Segoe UI"

        box = slide.shapes.add_textbox(Inches(x), Inches(1.55), Inches(4.3), Inches(5.2))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(citems):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = item
            p.font.size = Pt(15)
            p.font.name = "Segoe UI"
            p.font.color.rgb = INK
            p.space_after = Pt(8)

    add_footer(slide)


def database_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "6. Database Design")

    intro = slide.shapes.add_textbox(Inches(0.6), Inches(1.1), Inches(8.8), Inches(0.5))
    p = intro.text_frame.paragraphs[0]
    p.text = "PostgreSQL — 13 tables · React + Express three-tier architecture"
    p.font.size = Pt(14)
    p.font.color.rgb = MUTED
    p.font.name = "Segoe UI"

    core = [
        "users — Manager & Worker accounts (JWT login)",
        "products — product catalogue, shelf life, loss rate",
        "batches — retort date, printed vs correct expiry",
        "defects — main defect records & workflow status",
        "corrective_actions — assigned handling tasks",
        "root_cause_investigation — worker suspect, manager confirm",
        "evidence — photo uploads (defect & CA)",
        "activity_logs — workflow audit (13 action types)",
    ]
    add_bullets(slide, core, top=1.55, size=15)

    flow = slide.shapes.add_shape(1, Inches(0.6), Inches(5.85), Inches(8.8), Inches(0.75))
    flow.fill.solid()
    flow.fill.fore_color.rgb = LIGHT_BG
    flow.line.color.rgb = TEAL
    flow.line.width = Pt(1)

    ft = slide.shapes.add_textbox(Inches(0.75), Inches(5.95), Inches(8.5), Inches(0.55))
    p = ft.text_frame.paragraphs[0]
    p.text = "Core flow:  Product → Batch → Defect → Corrective Action → Root Cause → Close"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TEAL_DARK
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.CENTER

    note = slide.shapes.add_textbox(Inches(0.6), Inches(6.55), Inches(8.8), Inches(0.4))
    p = note.text_frame.paragraphs[0]
    p.text = "Batch no.: {product_code}-B-{YYYYMMDD}  ·  Insert ERD: diagrams/erd-qdts-core-report.svg"
    p.font.size = Pt(11)
    p.font.color.rgb = MUTED
    p.font.name = "Segoe UI"

    add_footer(slide)


def conclusion_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_header_bar(slide, "7. Conclusion")

    add_bullets(slide, [
        "Working QDTS prototype delivered for Kak Norie quality operations",
        "Replaces paper-based defect recording with structured digital workflow",
        "Expiry mismatch alerts & batch traceability address client-specific needs",
        "Role-based access (Manager / Worker), reports, CSV export, activity log",
        "Automated API tests pass (smoke + integration)",
        "Limitations: localhost prototype, in-app notifications only, no ERP/AI",
        "Next: production deployment, email alerts, mobile PWA",
        "",
        "→ Live system demonstration (~15 minutes)",
    ], top=1.2, size=17)

    cta = slide.shapes.add_shape(1, Inches(2.5), Inches(6.2), Inches(5), Inches(0.55))
    cta.fill.solid()
    cta.fill.fore_color.rgb = TEAL
    cta.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(2.5), Inches(6.28), Inches(5), Inches(0.4))
    p = tb.text_frame.paragraphs[0]
    p.text = "Demo: http://localhost:5173"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.CENTER

    add_footer(slide)


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    title_slide(prs)

    content_slide(prs, "1. Introduction", [
        "Client: Kak Norie / Retort Niaga — small-scale retort food manufacturing",
        "Products: shelf-stable pouches & bottles (200g–250g, ~12-month shelf life)",
        "System: QDTS — Quality Defect Tracking System (web application)",
        "Technology: React 19 frontend · Express 5 API · PostgreSQL database",
        "Users: Manager (nazhif) + Workers (siti_aminah, hairul_nizam)",
        "Purpose: digitise defect lifecycle from report → review → action → close",
    ])

    content_slide(prs, "2. Problem Statement", [
        "Defects recorded in paper notebooks — slow search, no central archive",
        "Corrective actions not monitored — same problems can repeat",
        "No formal batch numbers — weak link between defect and production batch",
        "Printed pouch expiry may differ from correct retort-based expiry",
        "Loss and disposal data scattered across paper notes",
        "Limited audit trail — no timestamps or user accountability",
        "Primary data: WhatsApp (9 Jun 2026) + Google Form (11 Jun 2026)",
    ])

    content_slide(prs, "3. Objectives", [
        "O1 — Structured digital defect records linked to product & batch",
        "O2 — Role-based workflow: report → review → assign → complete → verify → close",
        "O3 — Auto expiry validation & mismatch alerts",
        "O4 — Corrective action tracking with due dates, evidence & verification",
        "O5 — Product loss calculation & management reports",
        "O6 — CSV export & browser print summaries for audit",
        "O7 — Activity log for compliance and traceability",
    ])

    two_column_slide(prs, "4. Scope",
        "In Scope", [
            "Defect reporting (9 detection stages)",
            "Manager review & corrective action assignment",
            "Root cause suspect & manager confirmation",
            "Batch master data & expiry mismatch reports",
            "Dashboard KPIs, notifications, activity log",
            "CSV export & print summary",
            "Web browser on Windows PC (localhost demo)",
        ],
        "Out of Scope (v1)", [
            "AI / YOLO visual defect detection",
            "OCR label reading",
            "Blockchain traceability",
            "ERP integration",
            "Email / SMS alerts (in-app bell only)",
            "Native mobile app",
            "Full enterprise CAPA / QMS platform",
        ])

    content_slide(prs, "5. Methodology", [
        "Research approach (adapted from Bharosa et al., 2009):",
        "  1. Literature review — food quality, traceability, RBAC",
        "  2. Case study — Kak Norie manual as-is process",
        "  3. Requirement confirmation — WhatsApp first, then Google Form",
        "SDLC: Modified waterfall + Object-Oriented Analysis & Design (OOAD)",
        "Phases: Requirements → Analysis → Design → Implementation → Testing",
        "Tools: VS Code, Node.js, PostgreSQL, Git, Figma (diagrams)",
    ])

    database_slide(prs)
    conclusion_slide(prs)

    prs.save(OUT)
    print(f"Created: {OUT}")
    print(f"Slides: {len(prs.slides)} (1 title + 7 content)")


if __name__ == "__main__":
    main()
